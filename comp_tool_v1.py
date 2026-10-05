import difflib
import streamlit as st
from pypdf import PdfReader
from docx import Document
import io

# Configure the Streamlit page layout
st.set_page_config(
    page_title="Ultimate File Comparison Tool",
    page_icon="🔍",
    layout="wide"
)

st.title("🔍 Ultimate File Comparison Tool")
st.markdown("Upload **TXT, PDF, or Word (.docx)** files to highlight insertions, deletions, and line changes.")

def extract_text_from_file(uploaded_file) -> list:
    """Extracts text content from TXT, PDF, or DOCX files and returns a list of lines."""
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    
    # Reset stream pointer after reading raw bytes
    uploaded_file.seek(0)
    
    # Handle PDF Files
    if filename.endswith('.pdf'):
        pdf_stream = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_stream)
        text_lines = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_lines.extend(page_text.splitlines())
        return text_lines
        
    # Handle Word Documents (.docx)
    elif filename.endswith('.docx'):
        docx_stream = io.BytesIO(file_bytes)
        doc = Document(docx_stream)
        text_lines = [p.text for p in doc.paragraphs if p.text.strip() != ""]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text_lines.extend(cell.text.splitlines())
        return text_lines
        
    # Handle Plain Text / Source Code Files
    else:
        try:
            return file_bytes.decode("utf-8").splitlines()
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1").splitlines()

# Render two upload columns layout
col1, col2 = st.columns(2)

with col1:
    st.subheader("Original File (File A)")
    file_a = st.file_uploader("Choose original file", type=["txt", "pdf", "docx", "py", "js", "json", "csv"], key="file_a")

with col2:
    st.subheader("Modified File (File B)")
    file_b = st.file_uploader("Choose modified file", type=["txt", "pdf", "docx", "py", "js", "json", "csv"], key="file_b")

# Sidebar Configuration Settings
st.sidebar.header("🔧 Configuration Options")
show_mode = st.sidebar.radio("View Mode", ["Show All Lines", "Show Changes Only"])
ignore_whitespace = st.sidebar.checkbox("Ignore trailing whitespace", value=True)
wrap_lines = st.sidebar.checkbox("Enable line wrapping", value=True)

if file_a and file_b:
    try:
        # Extract lines
        lines_a = extract_text_from_file(file_a)
        lines_b = extract_text_from_file(file_b)

        if ignore_whitespace:
            lines_a = [line.rstrip() for line in lines_a if line.strip() != ""]
            lines_b = [line.rstrip() for line in lines_b if line.strip() != ""]

        if not lines_a and not lines_b:
            st.warning("Both files appear to be empty or unreadable.")
        else:
            st.success(f"Analyzing differences between '{file_a.name}' and '{file_b.name}'...")

            # Run diff generation engine
            differ = difflib.HtmlDiff(wrapcolumn=80 if wrap_lines else None)
            use_context = True if show_mode == "Show Changes Only" else False
            
            # Generate raw table
            raw_diff_table = differ.make_table(
                lines_a, 
                lines_b, 
                fromdesc=file_a.name, 
                todesc=file_b.name, 
                context=use_context,
                numlines=3
            )

            # FIXED: Build full standalone HTML document with embedded CSS inside the frame
            styled_html_output = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <meta charset="utf-8">
                <style>
                    body {{
                        margin: 0;
                        padding: 10px;
                        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
                        background-color: #ffffff;
                    }}
                    table.diff {{
                        font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, monospace !important;
                        font-size: 13px !important;
                        border-collapse: collapse !important;
                        width: 100% !important;
                        border: 1px solid #e5e7eb !important;
                    }}
                    table.diff td {{
                        padding: 4px 8px !important;
                        line-height: 1.5 !important;
                        vertical-align: top !important;
                        word-break: break-all !important;
                    }}
                    table.diff th {{
                        background-color: #f3f4f6 !important;
                        color: #374151 !important;
                        font-weight: 600 !important;
                        padding: 8px !important;
                        border-bottom: 2px solid #e5e7eb !important;
                        text-align: left;
                    }}
                    table.diff .diff_header {{
                        background-color: #f9fafb !important;
                        color: #9ca3af !important;
                        text-align: right !important;
                        width: 45px !important;
                        user-select: none !important;
                        border-right: 1px solid #e5e7eb !important;
                        font-weight: normal !important;
                    }}
                    /* Git styling classes mapped from Python's difflib output */
                    .diff_add {{ background-color: #e6ffec !important; color: #1b5e20 !important; }}
                    .diff_chg {{ background-color: #fef9c3 !important; color: #854d0e !important; }}
                    .diff_sub {{ background-color: #ffebe9 !important; color: #8a1e1b !important; }}
                    
                    /* Clean up empty gap line representations */
                    td.diff_empty {{ background-color: #f9fafb !important; }}
                </style>
            </head>
            <body>
                {raw_diff_table}
            </body>
            </html>
            """

            st.markdown("### 📊 Structural Comparison Results")
            # Render the complete styled document inside the Streamlit view element
            st.components.v1.html(styled_html_output, height=750, scrolling=True)

    except Exception as e:
        st.error(f"An unexpected parsing exception occurred: {str(e)}")
else:
    st.info("💡 Pro-Tip: Drop any combination of PDF, Word, or text files above to compare them.")
